"""Read-only DOCX table extraction for the formula reference content audit.

Run with the bundled document Python runtime. Emits JSON to stdout; never edits
the supplied document. Keeps every business, note and data-source paragraph.
"""
import hashlib
import json
import re
import sys
from pathlib import Path

from docx import Document


def extract(path: Path) -> dict:
    document = Document(path)
    if len(document.tables) != 1:
        raise ValueError("Expected the reviewed five-column formula table")
    records = []
    for row in document.tables[0].rows[1:]:
        if len(row.cells) != 5:
            raise ValueError("Unexpected source table layout")
        group, name, business, notes, source = row.cells
        match = re.fullmatch(r"(\d+\.\d+)\.\s*(.+)", name.text.strip())
        if match is None:
            raise ValueError(f"Unexpected indicator name: {name.text}")
        records.append({
            "id": match[1], "title": match[2], "group": group.text.strip(),
            "business": business.text.splitlines(),
            "notes": notes.text.splitlines(),
            "dataSources": source.text.splitlines(),
        })
    return {
        "name": path.name,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest().upper(),
        "records": records,
    }


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(extract(Path(sys.argv[1])), ensure_ascii=False, indent=2))
